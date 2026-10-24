import json
import re
from pathlib import Path

from ..providers import ModelRequest
from ..risk.engine import make_risk

RULES = [
    {
        "name": "undefined_error_flow",
        "area": "error_handling",
        "severity": "HIGH",
        "needles": ("error", "rechazo", "rechazada", "fallo", "timeout", "expira"),
        "title": "Flujo de error o excepción potencialmente incompleto",
        "hypothesis": "El requisito menciona una condición de error/excepción pero puede no definir completamente recuperación, mensaje o estado final.",
        "missing": ["resultado esperado", "mensaje al usuario", "estado final o recuperación"],
    },
    {
        "name": "boundary_without_behavior",
        "area": "business_rules",
        "severity": "MEDIUM",
        "needles": ("máximo", "maximo", "mínimo", "minimo", "límite", "limite", "hasta "),
        "title": "Regla de límite requiere comportamiento de borde",
        "hypothesis": "Existe un límite de negocio y conviene verificar qué ocurre debajo, exactamente en el límite y por encima.",
        "missing": ["comportamiento en el límite", "comportamiento fuera del límite"],
    },
    {
        "name": "conditional_rule",
        "area": "business_rules",
        "severity": "MEDIUM",
        "needles": ("si ", "cuando ", "en caso de"),
        "title": "Regla condicional requiere camino alternativo",
        "hypothesis": "Se define una condición y puede faltar el comportamiento cuando esa condición no se cumple.",
        "missing": ["camino alternativo", "estado esperado cuando la condición es falsa"],
    },
]

VISION_SYSTEM = """You are the Breakers visual QA analyst. Analyze only the supplied image from a functional document. Describe QA-relevant UI, diagrams, flows, labels, states, contradictions, missing paths or ambiguous behavior. Do not invent unreadable content. Return JSON with a single key named analysis containing concise factual text."""


def _sentences(text: str):
    return [x.strip() for x in re.split(r"(?<=[.!?])\s+|\n+", text) if x.strip()]


def _read_docx(path: Path):
    from docx import Document

    document = Document(path)
    chunks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                chunks.append(" | ".join(values))
    images = []
    for rel in document.part.rels.values():
        if "image" not in str(getattr(rel, "reltype", "")).lower():
            continue
        target_part = getattr(rel, "target_part", None)
        blob = getattr(target_part, "blob", None)
        if blob:
            images.append(bytes(blob))
    return "\n".join(chunks), len(document.tables), images


def _analyze_images(images: list[bytes], vision_provider):
    analyses = []
    failures = []
    if not images or vision_provider is None:
        return analyses, failures
    for index, image in enumerate(images, start=1):
        try:
            response = vision_provider.generate(ModelRequest(
                task=f"Analyze embedded document image {index} for QA-relevant evidence.",
                system=VISION_SYSTEM,
                images=[image],
            ))
            payload = json.loads(response.content or "{}")
            analysis = str(payload.get("analysis") or "").strip()
            if not analysis:
                raise ValueError("vision provider returned no analysis")
            analyses.append({
                "image": index,
                "analysis": analysis,
                "provider": response.provider,
                "model": response.model,
            })
        except (OSError, ValueError, json.JSONDecodeError) as exc:
            failures.append({"image": index, "error": str(exc)})
    return analyses, failures


def analyze_requirements_text(text: str, reference: str = "functional_document"):
    risks = []
    for index, sentence in enumerate(_sentences(text), start=1):
        lowered = sentence.lower()
        for rule in RULES:
            if any(needle in lowered for needle in rule["needles"]):
                risks.append(make_risk(
                    area=rule["area"], severity=rule["severity"], title=rule["title"],
                    hypothesis=rule["hypothesis"], source_type="functional_document",
                    reference=f"{reference}#fragment-{index}", excerpt=sentence[:500],
                    confidence=0.55, missing_evidence=rule["missing"],
                    suggested_action="Aclarar la regla antes de convertirla en desarrollo y pruebas.",
                ))
    return risks


def analyze_requirements_document(target: str, vision_provider=None):
    path = Path(target)
    suffix = path.suffix.lower()
    coverage = {
        "text_analyzed": False,
        "tables_processed": 0,
        "images_detected": 0,
        "images_analyzed": 0,
        "image_analyses": [],
        "image_failures": [],
        "status": "PARTIAL",
    }
    if suffix not in {".txt", ".md", ".docx"}:
        return [], ["El adaptador v0.3 analiza texto, Markdown y DOCX. PDF queda pendiente de extracción."], coverage
    try:
        if suffix == ".docx":
            text, table_count, images = _read_docx(path)
            coverage["tables_processed"] = table_count
            coverage["images_detected"] = len(images)
        else:
            text = path.read_text(encoding="utf-8")
            images = []
    except (OSError, UnicodeDecodeError, ValueError, KeyError):
        return [], ["No se pudo obtener contenido legible del documento."], coverage
    if not text.strip() and not images:
        return [], ["El documento no contiene contenido analizable."], coverage

    risks = analyze_requirements_text(text, path.name) if text.strip() else []
    coverage["text_analyzed"] = bool(text.strip())
    image_analyses, image_failures = _analyze_images(images, vision_provider)
    coverage["image_analyses"] = image_analyses
    coverage["image_failures"] = image_failures
    coverage["images_analyzed"] = len(image_analyses)
    coverage["status"] = "COMPLETE" if coverage["images_detected"] == coverage["images_analyzed"] else "PARTIAL"

    limitations = []
    if coverage["images_detected"] > coverage["images_analyzed"]:
        limitations.append(
            f'{coverage["images_detected"]} imágenes detectadas / {coverage["images_analyzed"]} analizadas. '
            "La cobertura documental es parcial hasta contar con análisis visual real."
        )
    return risks, limitations, coverage


def analyze_requirements_file(target: str):
    risks, limitations, _coverage = analyze_requirements_document(target)
    return risks, limitations

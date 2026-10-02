from __future__ import annotations

import csv
import io
from typing import Any

from openpyxl import Workbook


def build_strike_export(payload: dict[str, Any], fmt: str) -> tuple[bytes, str, str]:
    fmt = (fmt or "XLSX").upper()
    if fmt == "CSV":
        data = _csv_bytes(payload)
        return data, "text/csv; charset=utf-8", "BREAKERS-STRIKE-plan.csv"
    if fmt != "XLSX":
        raise ValueError("format must be XLSX or CSV")
    data = _xlsx_bytes(payload)
    return data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "BREAKERS-STRIKE-plan.xlsx"


def _tests(payload: dict[str, Any]) -> list[dict[str, Any]]:
    return list((payload.get("plan") or {}).get("tests") or payload.get("tests") or [])


def _test_row(test: dict[str, Any]) -> list[str]:
    return [
        str(test.get("id", "")),
        str(test.get("description", "")),
        " | ".join(map(str, test.get("preconditions") or [])),
        " | ".join(map(str, test.get("steps") or [])),
        str(test.get("expected_result", "")),
        str(test.get("origin", "")),
        str(test.get("executability", "")),
        ", ".join(map(str, test.get("coverage_unit_ids") or [])),
    ]


HEADERS = ["ID", "Descripción", "Precondiciones", "Pasos", "Resultado esperado", "Origen", "Ejecutabilidad", "Coverage Unit"]


def _csv_bytes(payload: dict[str, Any]) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(HEADERS)
    for test in _tests(payload):
        writer.writerow(_test_row(test))
    return output.getvalue().encode("utf-8-sig")


def _xlsx_bytes(payload: dict[str, Any]) -> bytes:
    workbook = Workbook()
    plan = workbook.active
    plan.title = "Plan de pruebas"
    plan.append(HEADERS)
    for test in _tests(payload):
        plan.append(_test_row(test))

    gaps = workbook.create_sheet("Definiciones pendientes")
    gaps.append(["ID", "Descripción", "Fuentes"])
    for gap in (payload.get("definition_gaps") or []):
        gaps.append([gap.get("id", ""), gap.get("description", ""), ", ".join(map(str, gap.get("source_refs") or []))])

    requirements = workbook.create_sheet("Requisitos para ejecutar")
    requirements.append(["ID", "Tipo", "Descripción", "Pruebas relacionadas"])
    execution = payload.get("execution") or {}
    for item in (execution.get("requirements") or payload.get("requirements") or []):
        requirements.append([item.get("id", ""), item.get("type", ""), item.get("description", ""), ", ".join(map(str, item.get("test_ids") or []))])

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()

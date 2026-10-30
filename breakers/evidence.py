from __future__ import annotations

from pathlib import Path
import re
from typing import Any


MAX_FRAGMENT_CHARS = 500
MAX_TOTAL_CHARS = 12_000
MAX_FRAGMENTS = 60

_SAFE_REPOSITORY_CONTEXT_FILES = {
    "readme",
    "readme.md",
    "readme.txt",
    "readme.rst",
    "pyproject.toml",
    "package.json",
    "requirements.txt",
    "dockerfile",
    "docker-compose.yml",
    "docker-compose.yaml",
    "compose.yml",
    "compose.yaml",
    "openapi.json",
    "openapi.yaml",
    "openapi.yml",
    "swagger.json",
    "swagger.yaml",
    "swagger.yml",
}

_CREDENTIAL_PATTERNS = (
    re.compile(r"(?i)(api[_-]?key|token|secret|password|passwd|pwd)(\s*[:=]\s*)([^\s,;]+)"),
    re.compile(r"(?i)(authorization\s*:\s*bearer\s+)([^\s,;]+)"),
    re.compile(r"(https?://)([^/@\s:]+):([^/@\s]+)@"),
)


def _sanitize_text(text: str) -> str:
    value = str(text or "")
    value = _CREDENTIAL_PATTERNS[0].sub(r"\1\2[REDACTED]", value)
    value = _CREDENTIAL_PATTERNS[1].sub(r"\1[REDACTED]", value)
    value = _CREDENTIAL_PATTERNS[2].sub(r"\1[REDACTED]@", value)
    return value


def _bounded_fragments(text: str, reference: str) -> list[dict[str, str]]:
    fragments: list[dict[str, str]] = []
    total = 0
    for index, raw in enumerate(text.replace("\r\n", "\n").split("\n"), start=1):
        value = _sanitize_text(raw.strip())
        if not value:
            continue
        value = value[:MAX_FRAGMENT_CHARS]
        remaining = MAX_TOTAL_CHARS - total
        if remaining <= 0:
            break
        value = value[:remaining]
        fragments.append({
            "kind": "semantic_evidence",
            "reference": f"{reference}#fragment-{index}",
            "text": value,
        })
        total += len(value)
        if len(fragments) >= MAX_FRAGMENTS:
            break
    return fragments


def _read_text_file(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return path.read_text(encoding="utf-8", errors="replace")


def _read_docx(path: Path) -> str:
    from docx import Document

    document = Document(path)
    chunks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                chunks.append(" | ".join(values))
    return "\n".join(chunks)


def _document_fragments(path: Path) -> list[dict[str, str]]:
    suffix = path.suffix.lower()
    try:
        if suffix in {".txt", ".md", ".json", ".yaml", ".yml", ".csv", ".log", ".xml", ".html"}:
            return _bounded_fragments(_read_text_file(path), path.name)
        if suffix == ".docx":
            return _bounded_fragments(_read_docx(path), path.name)
    except (OSError, ValueError, KeyError):
        return []
    return []


def _repository_fragments(path: Path) -> list[dict[str, str]]:
    fragments: list[dict[str, str]] = []
    selected: list[Path] = []

    try:
        for child in path.iterdir():
            if child.is_file() and child.name.lower() in _SAFE_REPOSITORY_CONTEXT_FILES:
                selected.append(child)
    except OSError:
        return []

    for candidate in sorted(selected, key=lambda item: item.name.lower()):
        try:
            candidate_fragments = _bounded_fragments(_read_text_file(candidate), candidate.name)
        except OSError:
            continue
        for fragment in candidate_fragments:
            if sum(len(item["text"]) for item in fragments) >= MAX_TOTAL_CHARS:
                return fragments
            fragments.append(fragment)
            if len(fragments) >= MAX_FRAGMENTS:
                return fragments
    return fragments


def collect_evidence_fragments(target: str, source: dict[str, Any]) -> list[dict[str, str]]:
    """Collect a bounded, redacted semantic sample for quality agents.

    Repository sampling deliberately avoids source files, .env files and credential-like
    artifacts. The purpose is to understand product/context signals, not to replace scanners.
    """
    path = Path(target)
    source_type = str(source.get("type") or "").strip().lower()
    input_type = str(source.get("input_type") or "").strip().lower()

    if (source_type == "repository" or input_type == "repository") and path.is_dir():
        return _repository_fragments(path)
    if path.is_file():
        return _document_fragments(path)
    return []

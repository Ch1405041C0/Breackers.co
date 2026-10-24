import json

from docx import Document
from docx.shared import Inches
from PIL import Image

from breakers.analyzers.requirements import analyze_requirements_document, analyze_requirements_file
from breakers.providers import ModelResponse


class FakeVisionProvider:
    def generate(self, request):
        assert request.images
        return ModelResponse(
            content=json.dumps({"analysis": "Pantalla con estado de error visible para el usuario."}),
            provider="fake-vision",
            model="vision-test",
        )


def test_docx_functional_document_is_extracted_and_analyzed(tmp_path):
    path = tmp_path / "consentimiento.docx"
    document = Document()
    document.add_paragraph("Cuando falla el servicio se debe mostrar un mensaje de error.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Regla"
    table.cell(0, 1).text = "En caso de timeout informar al usuario"
    document.save(path)

    risks, limitations = analyze_requirements_file(str(path))

    assert limitations == []
    assert risks
    assert any(item.area == "error_handling" for item in risks)


def test_docx_with_image_is_partial_until_image_is_really_analyzed(tmp_path):
    image_path = tmp_path / "screen.png"
    Image.new("RGB", (20, 20), "white").save(image_path)
    path = tmp_path / "with-image.docx"
    document = Document()
    document.add_paragraph("Cuando falla el servicio se debe mostrar un mensaje de error.")
    document.add_picture(str(image_path), width=Inches(1))
    document.save(path)

    risks, limitations, coverage = analyze_requirements_document(str(path))

    assert risks
    assert coverage["text_analyzed"] is True
    assert coverage["images_detected"] == 1
    assert coverage["images_analyzed"] == 0
    assert coverage["status"] == "PARTIAL"
    assert limitations


def test_docx_image_becomes_complete_only_after_real_vision_response(tmp_path):
    image_path = tmp_path / "screen.png"
    Image.new("RGB", (20, 20), "white").save(image_path)
    path = tmp_path / "with-vision.docx"
    document = Document()
    document.add_paragraph("Flujo funcional de consentimiento.")
    document.add_picture(str(image_path), width=Inches(1))
    document.save(path)

    _risks, limitations, coverage = analyze_requirements_document(
        str(path),
        vision_provider=FakeVisionProvider(),
    )

    assert limitations == []
    assert coverage["images_detected"] == 1
    assert coverage["images_analyzed"] == 1
    assert coverage["image_analyses"][0]["provider"] == "fake-vision"
    assert coverage["status"] == "COMPLETE"


def test_unsupported_document_reports_real_limitation(tmp_path):
    path = tmp_path / "requirements.pdf"
    path.write_bytes(b"not-a-real-pdf")
    risks, limitations = analyze_requirements_file(str(path))
    assert risks == []
    assert limitations
    assert "PDF" in limitations[0]
